"""Verificación catálogo tracción / dominadas / OLY-DLO vs Excel y API."""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.main import STRENGTH_EXERCISES, app, on_startup
from app.models import Athlete, Exercise
from app.strength_log_kinds import (
    HORIZONTAL_TRACTION_EXERCISE_NAME,
    HORIZONTAL_TRACTION_RM_COEFFICIENT,
    LOG_KIND_PULL_UP,
    PULL_UP_EXERCISE_NAMES,
    PULL_UP_MODALITIES,
    resolve_log_kind_for_exercise_name,
)
from app.strength_percentage import (
    EXERCISE_PERCENTAGE_CURVE_BY_NAME,
    build_percentage_table,
    resolve_percentage_curve,
)
from app.strength_rm import (
    EPLEY_OLY_DLO_COEFFICIENT,
    EXERCISE_RM_PROFILES,
    OLY_DLO_EXERCISE_NAMES,
    compute_estimated_rm,
    exercise_with_rm_profile,
)

EXCEL_PATH = Path(
    os.environ.get(
        "STRENGTH_EXCEL_PATH",
        r"c:\Users\santi\Downloads\Estimación 1RM - E-pley (4).xlsx",
    ),
)


@pytest.fixture(scope="module", autouse=True)
def _startup():
    on_startup()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def auth_headers(client: TestClient):
    res = client.post(
        "/auth/login",
        json={"email": "admin@ns.com", "password": os.environ["ADMIN_INITIAL_PASSWORD"]},
    )
    assert res.status_code == 200
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


class TestCatalogAlignment:
    def test_strength_exercises_count_24(self):
        assert len(STRENGTH_EXERCISES) == 24

    @pytest.mark.skipif(not EXCEL_PATH.is_file(), reason="Excel Nico no encontrado")
    def test_catalog_matches_excel_column_o(self):
        from openpyxl import load_workbook

        ws = load_workbook(EXCEL_PATH, read_only=True, data_only=True)["RM para app"]
        excel_names = []
        for row in range(35, 59):
            val = ws.cell(row, 15).value
            if val is not None and str(val).strip():
                excel_names.append(str(val).strip())
        assert set(excel_names) == set(STRENGTH_EXERCISES)

    def test_no_holly_hip_strap_in_catalog_yet(self):
        """Excel (4) no incluye Holly/Hip Strap; si Nico los agrega, actualizar catálogo."""
        joined = " ".join(STRENGTH_EXERCISES).lower()
        assert "holly" not in joined and "hollow" not in joined
        assert "hip strap" not in joined


class TestLogKinds:
    def test_all_dominadas_are_pull_up(self):
        for name in PULL_UP_EXERCISE_NAMES:
            assert name in STRENGTH_EXERCISES
            assert resolve_log_kind_for_exercise_name(name) == LOG_KIND_PULL_UP

    def test_horizontal_traction_is_rm_estimated(self):
        assert resolve_log_kind_for_exercise_name(HORIZONTAL_TRACTION_EXERCISE_NAME) == "rm_estimated"

    def test_db_exercises_log_kind_after_startup(self):
        with SessionLocal() as db:
            for name in PULL_UP_EXERCISE_NAMES:
                ex = db.query(Exercise).filter(Exercise.name == name).first()
                assert ex is not None, name
                assert ex.log_kind == LOG_KIND_PULL_UP


class TestHorizontalTractionRm:
    def test_epley_coefficient_0062(self):
        assert EXERCISE_RM_PROFILES[HORIZONTAL_TRACTION_EXERCISE_NAME] == (
            "epley",
            HORIZONTAL_TRACTION_RM_COEFFICIENT,
        )

    def test_percentage_curve_peso_muerto(self):
        ex = Exercise(name=HORIZONTAL_TRACTION_EXERCISE_NAME)
        ex = exercise_with_rm_profile(ex)
        assert resolve_percentage_curve(ex) == "peso_muerto"

    def test_rm_sample_100x5(self):
        ex = exercise_with_rm_profile(Exercise(name=HORIZONTAL_TRACTION_EXERCISE_NAME))
        # 100 * (1 + 5 * 0.062) = 131.0
        assert compute_estimated_rm(100.0, 5, ex) == pytest.approx(131.0)

    def test_percentage_table_matches_deadlift_curve(self):
        ex = Exercise(name=HORIZONTAL_TRACTION_EXERCISE_NAME)
        horizontal = build_percentage_table(100.0, exercise=ex)
        dead = build_percentage_table(100.0, curve_key="peso_muerto")
        assert horizontal == dead


class TestOlyDloEpleyProvisional:
    def test_all_oly_dlo_profiles_epley_033(self):
        for name in OLY_DLO_EXERCISE_NAMES:
            assert name in EXERCISE_RM_PROFILES
            assert EXERCISE_RM_PROFILES[name] == ("epley", EPLEY_OLY_DLO_COEFFICIENT)

    def test_clean_row_80x5_epley(self):
        ex = exercise_with_rm_profile(Exercise(name="Oly - Clean - Cargada"))
        assert compute_estimated_rm(80.0, 5, ex) == pytest.approx(80.0 * (1 + 5 * 0.033), abs=0.02)

    def test_snatch_70x2_epley(self):
        ex = exercise_with_rm_profile(Exercise(name="Oly - Snatch - Arranque"))
        assert compute_estimated_rm(70.0, 2, ex) == pytest.approx(70.0 * (1 + 2 * 0.033), abs=0.02)


class TestDominadasApi:
    def _athlete_and_exercise(self, dominada_name: str) -> tuple[int, int]:
        with SessionLocal() as db:
            athlete = db.query(Athlete).first()
            ex = db.query(Exercise).filter(Exercise.name == dominada_name).first()
            assert athlete is not None
            assert ex is not None
            return athlete.id, ex.id

    @pytest.mark.parametrize("dominada_name", sorted(PULL_UP_EXERCISE_NAMES))
    @pytest.mark.parametrize("modality,weight", [
        ("bodyweight", 0.0),
        ("band", 0.0),
        ("weighted", 10.0),
    ])
    def test_post_dominada_all_modalities(
        self,
        client: TestClient,
        auth_headers: dict,
        dominada_name: str,
        modality: str,
        weight: float,
    ):
        assert modality in PULL_UP_MODALITIES
        athlete_id, exercise_id = self._athlete_and_exercise(dominada_name)
        res = client.post(
            "/logs",
            headers=auth_headers,
            json={
                "athlete_id": athlete_id,
                "exercise_id": exercise_id,
                "date": "2026-10-08",
                "weight": weight,
                "reps": 8,
                "pull_up_modality": modality,
            },
        )
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["estimated_rm"] is None
        assert body["pull_up_modality"] == modality
        if modality == "weighted":
            assert body["weight"] == 10.0
        else:
            assert body["weight"] == 0.0

        summary = client.get(f"/logs/{body['id']}/summary", headers=auth_headers)
        assert summary.status_code == 404

        progress = client.get(
            f"/athletes/{athlete_id}/progress/{exercise_id}",
            headers=auth_headers,
        )
        assert progress.status_code == 200
        history = progress.json()["history"]
        assert any(
            h["reps"] == 8
            and h["pull_up_modality"] == modality
            and h["estimated_rm"] is None
            for h in history
        )


class TestHorizontalTractionApi:
    def test_post_remo_computes_rm(self, client: TestClient, auth_headers: dict):
        with SessionLocal() as db:
            athlete = db.query(Athlete).first()
            ex = (
                db.query(Exercise)
                .filter(Exercise.name == HORIZONTAL_TRACTION_EXERCISE_NAME)
                .first()
            )
            assert athlete and ex
            athlete_id, exercise_id = athlete.id, ex.id

        res = client.post(
            "/logs",
            headers=auth_headers,
            json={
                "athlete_id": athlete_id,
                "exercise_id": exercise_id,
                "date": "2026-10-08",
                "weight": 60.0,
                "reps": 6,
            },
        )
        assert res.status_code == 200
        rm = float(res.json()["estimated_rm"])
        expected = round(60.0 * (1 + 6 * HORIZONTAL_TRACTION_RM_COEFFICIENT), 2)
        assert abs(rm - expected) < 0.02

        log_id = res.json()["id"]
        summary = client.get(f"/logs/{log_id}/summary", headers=auth_headers)
        assert summary.status_code == 200
        assert summary.json()["percentage_curve"] == "peso_muerto"


class TestEveryCatalogExerciseHasCurveOrPullUp:
    def test_rm_exercises_have_percentage_mapping_or_default(self):
        for name in STRENGTH_EXERCISES:
            if name in PULL_UP_EXERCISE_NAMES:
                continue
            ex = Exercise(name=name)
            ex = exercise_with_rm_profile(ex)
            curve = resolve_percentage_curve(ex)
            assert curve in ("sentadilla", "peso_muerto", "banco_plano")
            if name in EXERCISE_PERCENTAGE_CURVE_BY_NAME:
                assert EXERCISE_PERCENTAGE_CURVE_BY_NAME[name] == curve
