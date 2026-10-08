"""Compara STRENGTH_EXERCISES con columna O de 'RM para app' (E-pley v4)."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

from app.main import STRENGTH_EXERCISES

XLSX = Path(r"c:\Users\santi\Downloads\Estimación 1RM - E-pley (4).xlsx")

NEW_TRACTION_ROWS = (45, 46, 47, 48, 49)  # O45–O49 en RM para app
ENVION_CANONICAL = "Oly - Clean & Jerk - Envión"


def _char_diff(a: str, b: str) -> list[str]:
    if a == b:
        return []
    lines: list[str] = []
    max_len = max(len(a), len(b))
    for i in range(max_len):
        ca = a[i] if i < len(a) else "∅"
        cb = b[i] if i < len(b) else "∅"
        if ca != cb:
            lines.append(f"    @{i}: excel U+{ord(ca) if ca != '∅' else 0:04X} {ca!r}  |  app U+{ord(cb) if cb != '∅' else 0:04X} {cb!r}")
    if len(a) != len(b):
        lines.append(f"    longitudes: excel={len(a)} app={len(b)}")
    return lines


def main() -> None:
    if not XLSX.is_file():
        print(f"Missing {XLSX}", file=sys.stderr)
        sys.exit(1)
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["RM para app"]
    excel_names: list[str] = []
    excel_raw_by_row: dict[int, str] = {}
    for row in range(35, 59):
        val = ws[f"O{row}"].value
        if val is None or str(val).strip() == "":
            continue
        excel_raw_by_row[row] = str(val)
        excel_names.append(str(val).strip())

    print(f"Excel: {XLSX.name} — {len(excel_names)} nombres (O35:O58, .strip())\n")

    print("=== 5 ejercicios nuevos (tracción / dominadas) — Excel vs app ===\n")
    print("| # | Excel (O) | App | ¿Igual? | Prefijo |")
    print("|---|-----------|-----|---------|---------|")
    all_new_ok = True
    for row in NEW_TRACTION_ROWS:
        excel_s = excel_raw_by_row[row].strip()
        app_s = STRENGTH_EXERCISES[row - 35]
        equal = excel_s == app_s
        if "Vertical" in excel_s:
            prefix_ok = app_s.startswith("Traccion Vertical - ")
        else:
            prefix_ok = app_s.startswith("Tracción Horizontal - ")
        if not equal or not prefix_ok:
            all_new_ok = False
        print(f"| O{row} | {excel_s} | {app_s} | {'OK' if equal else 'DIFF'} | {'OK' if prefix_ok else 'FAIL'} |")
        if not equal:
            for line in _char_diff(excel_s, app_s):
                print(line)
    print()

    print("=== Regresión 19 ejercicios (sin bloque tracción O45–O49) ===")
    legacy_excel = [n for i, n in enumerate(excel_names) if (35 + i) not in NEW_TRACTION_ROWS]
    legacy_app = [
        n for i, n in enumerate(STRENGTH_EXERCISES) if (35 + i) not in NEW_TRACTION_ROWS
    ]
    legacy_ok = legacy_excel == legacy_app
    print(f"  Orden y texto: {'OK' if legacy_ok else 'FALLO'} ({len(legacy_excel)} ítems)")
    if not legacy_ok:
        for i, (a, b) in enumerate(zip(legacy_excel, legacy_app, strict=False)):
            if a != b:
                print(f"  @{i}: excel={a!r} app={b!r}")
    print()

    raw_envion = excel_raw_by_row.get(53, "")
    if raw_envion.strip() == ENVION_CANONICAL and raw_envion != ENVION_CANONICAL:
        print(
            f"Nota Envión: Excel raw len={len(raw_envion)} (espacio final); "
            f"app usa {ENVION_CANONICAL!r} len={len(ENVION_CANONICAL)}"
        )
    print()

    app_set = set(STRENGTH_EXERCISES)
    excel_set = set(excel_names)
    only_excel = excel_set - app_set
    only_app = app_set - excel_set
    if only_excel:
        print("Solo en Excel:", only_excel)
    if only_app:
        print("Solo en app:", only_app)

    order_ok = excel_names == STRENGTH_EXERCISES
    if order_ok and not only_excel and not only_app and all_new_ok and legacy_ok:
        print("OK: 24/24 coincidencia exacta (orden + conjunto + 5 nuevos + regresión 19)")
        return

    if not order_ok:
        for i, (a, b) in enumerate(zip(excel_names, STRENGTH_EXERCISES, strict=False)):
            if a != b:
                print(f"Orden diff @{i}: excel={a!r} app={b!r}")
    sys.exit(1)


if __name__ == "__main__":
    main()
